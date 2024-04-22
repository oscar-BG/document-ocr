from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.conf import settings
from django.core.files import File
from openpyxl import Workbook
from .forms import SubirDumentoImagenForm
from .models import SaveFileCsv
from google.cloud import documentai
from google.api_core.client_options import ClientOptions
from dotenv import load_dotenv
import os
import re
import locale
import json


os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = 'media/credencials/ocr-document.json'
load_dotenv()
locale.setlocale(locale.LC_ALL, 'es_ES.UTF-8')


endpoint        = os.getenv('ENDPOINT')
location        = os.getenv('LOCATION') 
project_id      = os.getenv('PROJECT_ID')
processor_id    = os.getenv('PROCESSOR_ID')

# Create your views here.

def formDocument(request):
    return render(request, 'index.html')

def homepage(request):
    form = SubirDumentoImagenForm()
    return render(request, 'index.html', {'form': form})


def upload(request):
    if request.method == "POST":
        form = SubirDumentoImagenForm(request.POST, request.FILES)
        if form.is_valid():
            instancia = form.save(commit=False)
            instancia.save()
            url_document = instancia.documento.url
            url_format = url_document[1:]
            get_text_form_pdf_ocr(url_format)

        return JsonResponse({"success" : True})


def listarData(request):
    # return render(request, 'show_csv.html')

    data = SaveFileCsv.objects.all()
    print(data)
    return render(request=request, template_name="show_csv.html", context={'data' : data})
    
def get_text_form_pdf_ocr(file_path):
    try:
        mime_type = 'application/pdf'
        client = documentai.DocumentProcessorServiceClient(client_options=ClientOptions(api_endpoint=f"{location}-{endpoint}"))

        name = client.processor_path(project_id, location, processor_id)
        with open(file_path, "rb") as image:
            image_content = image.read()

        raw_document = documentai.RawDocument(
            content=image_content, mime_type=mime_type)
        
        request = documentai.ProcessRequest(name=name, raw_document=raw_document)
        response = client.process_document(request=request)

        document = response.document
        data = []
        

        print(f"There are {len(document.pages)} page(s) in this document.")
        # Convertir el objeto document a un diccionario
        # serialized_document = documentai.Document.to_dict(document)
        # Guardar el diccionario como JSON
        # with open("media/documentPag15.json", "w") as json_file:
            # json.dump(serialized_document, json_file, indent=4)
        
        

        
        array_card_id = {}
        array_full_name = {}
        for entity in document.entities:
            card_id = ""
            full_name = ""
            pagina = int(entity.page_anchor.page_refs[0].page) + 1

            if entity.type_ == 'CARD_ID':
                if re.match('^[0-9]+$', entity.mention_text):
                    card_id = int(entity.mention_text)
                else:
                    texto_limpio = re.sub('[^0-9]', '', entity.mention_text)
                    card_id = int(texto_limpio)
                if pagina not in array_card_id:
                    array_card_id[pagina] = []
                array_card_id[pagina].append(card_id)
            elif entity.type_ == 'FULL_NAME':
                full_name = entity.mention_text
                full_name = full_name.replace("\n", " ")
                if pagina not in array_full_name:
                    array_full_name[pagina] = []
                array_full_name[pagina].append(full_name)

        
        # Ordenar las listas de cada página
        for page in array_card_id:
            array_card_id[page] = sorted(array_card_id[page])
        for page in array_full_name:
            array_full_name[page] = sorted(array_full_name[page], key=locale.strxfrm)


        print(page)

        document_file_name = os.path.splitext(os.path.basename(file_path))[0] + ".xlsx"
        document_file_path = os.path.join("media/csv/", document_file_name)
        relative_document_file_path = os.path.relpath(document_file_path, settings.MEDIA_ROOT)

        # Crear libro
        workbook = Workbook()

        # Escribir datos en hojas separadas por página
        for page, card_ids in array_card_id.items():
            hoja = workbook.create_sheet(title=f'Página {page}')

            hoja.append(['CARD_ID', 'FULL_NAME'])
            full_names = array_full_name.get(page, [])

            for card_id, full_name in zip(card_ids, full_names):
                hoja.append([card_id, full_name])

        # Eliminar la hoja predeterminada
        workbook.remove(workbook['Sheet'])

        # Guardar el libro de trabajo como un archivo Excel
        workbook.save(filename=document_file_path)

        new_csv_instance = SaveFileCsv.objects.create(excel_file=relative_document_file_path, name_document=document_file_name)
        
        return True
    except Exception as e:
        print(e)
        return None
    
